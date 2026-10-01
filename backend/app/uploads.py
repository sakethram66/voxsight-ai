import base64
import io
import uuid

from .config import Settings
from .session import Attachment

MAX_DOC_CHARS = 150_000
TEXT_EXT = {".txt", ".md", ".log", ".csv", ".json", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c",
            ".cpp", ".h", ".cs", ".go", ".rs", ".rb", ".php", ".html", ".css", ".yaml", ".yml", ".xml", ".sql", ".sh"}


class UploadError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status, self.message = status, message


def sniff_image(data: bytes):
    if data.startswith(b"\x89PNG\r\n\x1a\n"): return "image/png"
    if data.startswith(b"\xff\xd8\xff"): return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"): return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP": return "image/webp"
    return None


def _docx_text(data: bytes) -> str:
    from docx import Document
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _limit(data: bytes, mb: int, what: str):
    if len(data) > mb * 1024 * 1024:
        raise UploadError(413, f"{what} is too large (max {mb} MB).")


def classify(name: str, data: bytes, s: Settings, tag: str = "upload") -> Attachment:
    if not data:
        raise UploadError(400, "The file is empty.")
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    aid = uuid.uuid4().hex[:10]
    mime = sniff_image(data)
    if mime:
        _limit(data, s.max_image_mb, "Image")
        return Attachment(aid, name, mime, "image", data=data, tag=tag)
    if data.startswith(b"%PDF-"):
        _limit(data, s.max_pdf_mb, "PDF")
        return Attachment(aid, name, "application/pdf", "pdf", data=data)
    if ext in TEXT_EXT:
        _limit(data, s.max_text_mb, "Text file")
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise UploadError(415, "That file isn't valid UTF-8 text.")
        if len(text) > MAX_DOC_CHARS:
            text = text[:MAX_DOC_CHARS] + "\n[...truncated]"
        return Attachment(aid, name, "text/plain", "text", text=text)
    if ext == ".docx" and data[:2] == b"PK":
        _limit(data, s.max_docx_mb, "Word document")
        try:
            text = _docx_text(data)
        except Exception:
            raise UploadError(422, "Couldn't read that .docx file.")
        if not text.strip():
            raise UploadError(422, "That Word document has no readable text.")
        return Attachment(aid, name, "text/plain", "text", text=text[:MAX_DOC_CHARS])
    raise UploadError(415, "Unsupported file type. Use an image, PDF, .docx, or a text/code file.")


def frame_to_attachment(f: dict, s: Settings) -> Attachment:
    """Validate a live camera/screen frame arriving over the WebSocket."""
    tag = f.get("tag") if isinstance(f, dict) else None
    if tag not in ("camera", "screen"):
        raise UploadError(400, "Invalid live frame.")
    try:
        data = base64.b64decode(f.get("data") or "", validate=True)
    except Exception:
        raise UploadError(400, "Live frame was not valid base64.")
    mime = sniff_image(data)
    if not mime or len(data) > 4 * 1024 * 1024:
        raise UploadError(400, "Live frame must be an image under 4 MB.")
    label = "Camera snapshot (live)" if tag == "camera" else "Screen capture (live)"
    return Attachment(uuid.uuid4().hex[:10], label, mime, "image", data=data, tag=tag)
